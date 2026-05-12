"""Pre-deployment regression checks for security and stability fixes."""

from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from data_preprocessor import DataPreprocessor
from preprocessing_portal.server import (
    MAX_REQUEST_BYTES,
    _normalize_prefix,
    _resolve_prefixes,
)


class DataPreprocessorRegressionTests(unittest.TestCase):
    def test_load_dataframe_clears_stale_date_metadata(self) -> None:
        preprocessor = DataPreprocessor()
        ok, _ = preprocessor.load_dataframe(
            pd.DataFrame({"Date": ["2026-01-01 00:00:00"], "Value": [1]})
        )
        self.assertTrue(ok)
        self.assertEqual(preprocessor.date_column, "Date")

        ok, _ = preprocessor.load_dataframe(pd.DataFrame({"a": [1], "b": [2]}))

        self.assertTrue(ok)
        self.assertIsNone(preprocessor.date_column)
        self.assertIsNone(preprocessor.original_date_format)

    def test_save_dataframe_escapes_formula_like_strings_for_xlsx_and_csv(self) -> None:
        preprocessor = DataPreprocessor()
        df = pd.DataFrame(
            {
                '=HYPERLINK("http://evil-header")': ["=HYPERLINK(\"http://evil\")"],
                "+SUM(1,2)": ["+SUM(1,2)"],
                "\r=CMD": ["\n=PAYLOAD"],
                "safe": ["safe"],
            }
        )
        ok, _ = preprocessor.load_dataframe(df)
        self.assertTrue(ok)

        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "out.xlsx"
            csv_path = Path(tmpdir) / "out.csv"
            preprocessor._save_dataframe(preprocessor.processed_df, xlsx_path)
            preprocessor._save_dataframe(preprocessor.processed_df, csv_path)

            wb = load_workbook(xlsx_path, data_only=False)
            ws = wb.active
            self.assertEqual(ws["A1"].value, "'=HYPERLINK(\"http://evil-header\")")
            self.assertEqual(ws["B1"].value, "'+SUM(1,2)")
            self.assertEqual(ws["C1"].value, "'=CMD")
            self.assertEqual(ws["A2"].value, "'=HYPERLINK(\"http://evil\")")
            self.assertEqual(ws["B2"].value, "'+SUM(1,2)")
            self.assertEqual(ws["C2"].value, "'=PAYLOAD")
            self.assertNotEqual(ws["A1"].data_type, "f")

            with csv_path.open("r", encoding="utf-8-sig", newline="") as f:
                rows = list(csv.reader(f))
            self.assertEqual(rows[0][0], "'=HYPERLINK(\"http://evil-header\")")
            self.assertEqual(rows[0][1], "'+SUM(1,2)")
            self.assertEqual(rows[0][2], "'=CMD")
            self.assertEqual(rows[1][0], "'=HYPERLINK(\"http://evil\")")
            self.assertEqual(rows[1][1], "'+SUM(1,2)")
            self.assertEqual(rows[1][2], "'=PAYLOAD")

    def test_load_data_drops_unnamed_columns_without_crashing(self) -> None:
        preprocessor = DataPreprocessor()
        with tempfile.TemporaryDirectory() as tmpdir:
            csv_path = Path(tmpdir) / "with_unnamed.csv"
            csv_path.write_text("Unnamed: 0,Date,Value\n0,2026-01-01 00:00:00,10\n", encoding="utf-8")
            ok, message = preprocessor.load_data(str(csv_path))

        self.assertTrue(ok, message)
        self.assertNotIn("Unnamed: 0", preprocessor.columns)
        self.assertEqual(preprocessor.columns, ["Date", "Value"])

    def test_removed_source_indices_survive_index_reset(self) -> None:
        preprocessor = DataPreprocessor()
        ok, _ = preprocessor.load_dataframe(
            pd.DataFrame(
                {
                    "Date": pd.date_range("2026-01-01", periods=5, freq="min"),
                    "Value": [10, 20, 30, 40, 50],
                }
            )
        )
        self.assertTrue(ok)
        ok, _ = preprocessor.apply_filters(
            [{"column": "Value", "operator": ">=", "value": 30}]
        )
        self.assertTrue(ok)
        self.assertEqual(preprocessor.processed_df.index.tolist(), [0, 1, 2])
        self.assertEqual(preprocessor.get_removed_source_indices(), [0, 1])


class PortalSecurityRegressionTests(unittest.TestCase):
    def test_prefix_normalization_blocks_traversal(self) -> None:
        self.assertEqual(_normalize_prefix("pj1"), "PJ1")
        self.assertIsNone(_normalize_prefix("../PJ1"))
        self.assertIsNone(_normalize_prefix("PJ1/../../secret"))
        self.assertEqual(_resolve_prefixes({"prefix": ["../PJ1", "KY"]}), ("KY",))

    def test_request_body_limit_constant_is_bounded(self) -> None:
        self.assertLessEqual(MAX_REQUEST_BYTES, 1_000_000)


if __name__ == "__main__":
    unittest.main()
